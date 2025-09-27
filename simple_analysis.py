import openpyxl

# Load the Excel file
wb = openpyxl.load_workbook('order_type.xlsx')
ws = wb.active

print('Sheet name:', ws.title)
print('Max row:', ws.max_row)
print('Max column:', ws.max_column)

print('\nHeaders:')
for col in range(1, min(ws.max_column+1, 6)):
    print(f'  Column {col}: {ws.cell(1, col).value}')

print('\nSample data (first 5 rows):')
for row in range(1, min(6, ws.max_row+1)):
    row_data = []
    for col in range(1, min(ws.max_column+1, 6)):
        row_data.append(ws.cell(row, col).value)
    print(f'  Row {row}: {row_data}')

# Convert to CSV
print('\nConverting to CSV...')
with open('order_data.csv', 'w', encoding='utf-8') as f:
    for row in range(1, ws.max_row + 1):
        row_data = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row, col).value
            if cell_value is None:
                cell_value = ''
            row_data.append(str(cell_value))
        f.write(','.join(row_data) + '\n')

print('✅ Data converted to order_data.csv')
